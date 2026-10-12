import os
import re
import hashlib
from typing import Tuple

import docx
from docx.oxml.ns import qn

from src.schemas import (
    NormalizedDocument,
    DocumentMetadata,
    Section,
    TextBlock,
    TableBlock,
    ListBlock,
    ListItem,
)


# ============================================================
# TEXT EXTRACTION
# ============================================================

def get_paragraph_text(paragraph) -> str:
    """Extract visible text including hyperlink text in XML order."""
    return "".join(node.text for node in paragraph._element.iter(qn("w:t")) if node.text).strip()


# ============================================================
# WORD NUMBERING
# ============================================================

def build_numbering_map(doc):
    """Resolve numId -> abstractNumId -> ilvl -> numFmt."""
    result = {}

    try:
        root = doc.part.numbering_part.element
    except Exception:
        return result

    abstract_levels = {}

    for abstract_num in root.findall(qn("w:abstractNum")):
        abstract_id = abstract_num.get(qn("w:abstractNumId"))

        if abstract_id is None:
            continue

        level_formats = {}

        for lvl in abstract_num.findall(qn("w:lvl")):
            ilvl = lvl.get(qn("w:ilvl"))

            if ilvl is None:
                continue

            num_fmt = lvl.find(qn("w:numFmt"))
            fmt = num_fmt.get(qn("w:val")) if num_fmt is not None else "unknown"

            try:
                level_formats[int(ilvl)] = fmt
            except ValueError:
                continue

        abstract_levels[str(abstract_id)] = level_formats

    for num in root.findall(qn("w:num")):
        num_id = num.get(qn("w:numId"))
        abstract_ref = num.find(qn("w:abstractNumId"))

        if num_id is None or abstract_ref is None:
            continue

        abstract_id = abstract_ref.get(qn("w:val"))

        try:
            numeric_num_id = int(num_id)
        except (ValueError, TypeError):
            continue

        result[numeric_num_id] = abstract_levels.get(str(abstract_id), {})

    return result


def get_list_info(paragraph, doc, numbering_map=None) -> Tuple[bool, int, str]:
    """Return (is_list, ilvl, list_type)."""
    p_pr = paragraph._p.pPr

    if p_pr is None or p_pr.numPr is None or p_pr.numPr.numId is None:
        return False, 0, ""

    num_pr = p_pr.numPr

    try:
        num_id = int(num_pr.numId.val)
    except Exception:
        return False, 0, ""

    try:
        ilvl = int(num_pr.ilvl.val) if num_pr.ilvl is not None else 0
    except Exception:
        ilvl = 0

    if numbering_map is None:
        numbering_map = build_numbering_map(doc)

    fmt = numbering_map.get(num_id, {}).get(ilvl, "unknown")

    if fmt == "bullet":
        return True, ilvl, "unordered"

    if fmt in {"decimal", "lowerLetter", "upperLetter", "lowerRoman", "upperRoman"}:
        return True, ilvl, "ordered"

    style_name = paragraph.style.name if paragraph.style else ""

    if "Bullet" in style_name:
        return True, ilvl, "unordered"

    return True, ilvl, "ordered"


# ============================================================
# ORIGINAL DOCUMENT ORDER
# ============================================================

def iter_block_items(parent):
    """Yield paragraphs/tables in their real DOCX body order."""
    if not isinstance(parent, docx.document.Document):
        raise ValueError("Unsupported parent type")

    for child in parent.element.body.iterchildren():
        if child.tag == qn("w:p"):
            yield docx.text.paragraph.Paragraph(child, parent)

        elif child.tag == qn("w:tbl"):
            yield docx.table.Table(child, parent)


# ============================================================
# DOCX -> NORMALIZED DOCUMENT
# ============================================================

def convert_docx_to_normalized_document(file_path: str) -> NormalizedDocument:
    doc = docx.Document(file_path)
    numbering_map = build_numbering_map(doc)

    # --------------------------------------------------------
    # Enterprise
    # --------------------------------------------------------

    path_lower = file_path.lower()

    if "aventro" in path_lower:
        enterprise = "Aventro Motors"
    elif "cendara" in path_lower:
        enterprise = "Cendara University"
    elif "cloudway" in path_lower:
        enterprise = "CloudWay-24"
    elif "velvera" in path_lower:
        enterprise = "Velvera Technologies"
    elif "zx bank" in path_lower or "zx_bank" in path_lower or "zx-bank" in path_lower or "zx" in path_lower:
        enterprise = "ZX Bank"
    else:
        enterprise = "Unknown"

    identity = enterprise + "::" + os.path.basename(file_path)
    document_id = hashlib.sha256(identity.encode("utf-8")).hexdigest()

    title = None
    sections = []
    current_hierarchy = []
    current_blocks = []
    buffered_list_items = []

    # ========================================================
    # LIST BUILDER
    # ========================================================

    def build_list_block(items):
        if not items:
            return None

        base_level = items[0][1]
        root_type = items[0][2]
        root_block = ListBlock(list_type=root_type, items=[], page_number=None)

        last_item_at_level = {}

        for text, ilvl, list_type in items:
            level = max(base_level, ilvl)
            new_item = ListItem(text=text, sub_lists=[])

            if level == base_level:
                root_block.items.append(new_item)
                last_item_at_level[base_level] = new_item

                for old_level in list(last_item_at_level):
                    if old_level > base_level:
                        del last_item_at_level[old_level]

                continue

            parent_level = level - 1

            while parent_level >= base_level and parent_level not in last_item_at_level:
                parent_level -= 1

            if parent_level < base_level:
                root_block.items.append(new_item)
                last_item_at_level[base_level] = new_item
                continue

            parent_item = last_item_at_level[parent_level]

            if not parent_item.sub_lists or parent_item.sub_lists[-1].list_type != list_type:
                parent_item.sub_lists.append(ListBlock(list_type=list_type, items=[], page_number=None))

            parent_item.sub_lists[-1].items.append(new_item)
            last_item_at_level[level] = new_item

            for old_level in list(last_item_at_level):
                if old_level > level:
                    del last_item_at_level[old_level]

        return root_block

    def flush_list():
        if not buffered_list_items:
            return

        list_block = build_list_block(buffered_list_items)

        if list_block is not None and list_block.items:
            current_blocks.append(list_block)

        buffered_list_items.clear()

    # ========================================================
    # SECTION BUILDER
    # ========================================================

    def flush_section(force=False):
        flush_list()

        if current_blocks or (force and current_hierarchy):
            sections.append(Section(hierarchy=list(current_hierarchy), blocks=list(current_blocks)))
            current_blocks.clear()

    manual_number_re = re.compile(r"^(\d+)\.\s+(.+)$")

    # ========================================================
    # DOCUMENT EXTRACTION
    # ========================================================

    for block in iter_block_items(doc):

        # ----------------------------------------------------
        # PARAGRAPH
        # ----------------------------------------------------

        if isinstance(block, docx.text.paragraph.Paragraph):
            text = get_paragraph_text(block)

            if not text:
                continue

            style_name = block.style.name if block.style else ""

            # ------------------------------------------------
            # Heading
            # ------------------------------------------------

            if style_name.startswith("Heading"):
                try:
                    level = int(style_name.split()[-1])
                except (ValueError, IndexError):
                    level = 1

                level = max(1, min(level, 4))

                # If the incoming heading replaces the current terminal
                # heading, preserve the old heading-only section.
                force_heading_section = bool(current_hierarchy) and level <= len(current_hierarchy)

                flush_section(force=force_heading_section)

                if title is None and level == 1:
                    title = text

                current_hierarchy = current_hierarchy[:level - 1]
                current_hierarchy.append(text)
                continue

            # ------------------------------------------------
            # Title style
            # ------------------------------------------------

            if style_name == "Title":
                flush_section(force=bool(current_hierarchy))

                if title is None:
                    title = text

                current_hierarchy = [text]
                continue

            # ------------------------------------------------
            # Native Word list
            # ------------------------------------------------

            is_list, ilvl, list_type = get_list_info(block, doc, numbering_map)

            if is_list:
                if buffered_list_items:
                    base_level = buffered_list_items[0][1]
                    previous_type = buffered_list_items[-1][2]

                    if ilvl <= base_level and previous_type != list_type:
                        flush_list()

                buffered_list_items.append((text, ilvl, list_type))
                continue

            # ------------------------------------------------
            # Manual 1. 2. 3. list
            # ------------------------------------------------

            manual_match = manual_number_re.match(text)

            if manual_match:
                clean_text = manual_match.group(2).strip()

                if buffered_list_items and buffered_list_items[-1][2] != "ordered":
                    flush_list()

                buffered_list_items.append((clean_text, 0, "ordered"))
                continue

            # ------------------------------------------------
            # Normal text
            # ------------------------------------------------

            flush_list()
            current_blocks.append(TextBlock(text=text, page_number=None))

        # ----------------------------------------------------
        # TABLE
        # ----------------------------------------------------

        elif isinstance(block, docx.table.Table):
            flush_list()
            rows_data = []

            for row in block.rows:
                row_cells = []

                for cell in row.cells:
                    cell_parts = []

                    for paragraph in cell.paragraphs:
                        paragraph_text = get_paragraph_text(paragraph)

                        if paragraph_text:
                            cell_parts.append(paragraph_text)

                    row_cells.append("\n".join(cell_parts))

                rows_data.append(row_cells)

            if not rows_data:
                continue

            if not any(cell.strip() for row in rows_data for cell in row):
                continue

            first_row = rows_data[0]
            headers = []
            final_rows = rows_data

            non_empty_count = sum(1 for cell in first_row if cell.strip())
            bold_count = 0

            for index, cell in enumerate(block.rows[0].cells):
                if index >= len(first_row) or not first_row[index].strip():
                    continue

                if any(run.bold is True for paragraph in cell.paragraphs for run in paragraph.runs):
                    bold_count += 1

            if non_empty_count > 0 and bold_count / non_empty_count >= 0.5:
                headers = first_row
                final_rows = rows_data[1:]

            current_blocks.append(TableBlock(headers=headers, rows=final_rows, page_number=None))

    # Preserve final heading even if it has no body content.
    flush_section(force=True)

    if title is None:
        title = os.path.splitext(os.path.basename(file_path))[0]

    metadata = DocumentMetadata(
        enterprise=enterprise,
        source_filename=os.path.basename(file_path),
        source_format="docx",
        document_version=None,
        author=None,
        date=None,
        access_groups=[],
    )

    if not sections:
        sections = [Section(hierarchy=[], blocks=[])]

    return NormalizedDocument(
        document_id=document_id,
        metadata=metadata,
        title=title,
        sections=sections,
    )