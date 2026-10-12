import os
import re
import hashlib
import pptx

from pptx.enum.shapes import PP_PLACEHOLDER
from pptx.oxml.ns import qn

from src.schemas import NormalizedDocument, DocumentMetadata, Section, TextBlock, TableBlock, ListBlock, ListItem


# ============================================================
# TEXT
# ============================================================

def get_paragraph_text(paragraph):
    return paragraph.text.strip()


def get_shape_text(shape):
    if not shape.has_text_frame: return ""
    return "\n".join(get_paragraph_text(p) for p in shape.text_frame.paragraphs if get_paragraph_text(p)).strip()


# ============================================================
# BULLET / NUMBERING
# ============================================================

def bullet_type_from_ppr(ppr):
    if ppr is None: return None
    if ppr.find(qn("a:buNone")) is not None: return "none"
    if ppr.find(qn("a:buAutoNum")) is not None: return "ordered"
    if ppr.find(qn("a:buChar")) is not None or ppr.find(qn("a:buBlip")) is not None: return "unordered"
    return None


def get_shape_level_style(shape, level):
    try:
        nodes = shape._element.xpath(f".//a:lstStyle/a:lvl{level + 1}pPr")
        return nodes[0] if nodes else None
    except Exception:
        return None


def get_layout_placeholder(slide, shape):
    if not shape.is_placeholder: return None

    try:
        idx = shape.placeholder_format.idx
        for placeholder in slide.slide_layout.placeholders:
            if placeholder.placeholder_format.idx == idx: return placeholder
    except Exception:
        pass

    return None


def get_master_level_style(slide, shape, level):
    try:
        master = slide.slide_layout.slide_master
    except Exception:
        return None

    style_name = "otherStyle"

    if shape.is_placeholder:
        try:
            ph_type = shape.placeholder_format.type
            if ph_type in (PP_PLACEHOLDER.TITLE, PP_PLACEHOLDER.CENTER_TITLE): style_name = "titleStyle"
            elif ph_type in (PP_PLACEHOLDER.BODY, PP_PLACEHOLDER.OBJECT, PP_PLACEHOLDER.SUBTITLE): style_name = "bodyStyle"
        except Exception:
            pass

    try:
        nodes = master.element.xpath(f".//p:txStyles/p:{style_name}/a:lvl{level + 1}pPr")
        return nodes[0] if nodes else None
    except Exception:
        return None


def get_list_info(paragraph, shape, slide):
    level = paragraph.level or 0

    result = bullet_type_from_ppr(paragraph._p.pPr)
    if result == "none": return False, level, ""
    if result in ("ordered", "unordered"): return True, level, result

    result = bullet_type_from_ppr(get_shape_level_style(shape, level))
    if result == "none": return False, level, ""
    if result in ("ordered", "unordered"): return True, level, result

    layout_placeholder = get_layout_placeholder(slide, shape)

    if layout_placeholder is not None:
        result = bullet_type_from_ppr(get_shape_level_style(layout_placeholder, level))
        if result == "none": return False, level, ""
        if result in ("ordered", "unordered"): return True, level, result

    result = bullet_type_from_ppr(get_master_level_style(slide, shape, level))
    if result == "none": return False, level, ""
    if result in ("ordered", "unordered"): return True, level, result

    return False, level, ""


# ============================================================
# LIST BUILDER
# ============================================================

def build_list_block(items, slide_number):
    if not items: return None

    base_level = items[0][1]
    root = ListBlock(list_type=items[0][2], items=[], slide_number=slide_number)
    last_item_at_level = {}

    for text, ilvl, list_type in items:
        level = max(base_level, ilvl)
        item = ListItem(text=text, sub_lists=[])

        if level == base_level:
            root.items.append(item)
            last_item_at_level[base_level] = item
            for old_level in list(last_item_at_level):
                if old_level > base_level: del last_item_at_level[old_level]
            continue

        parent_level = level - 1
        while parent_level >= base_level and parent_level not in last_item_at_level: parent_level -= 1

        if parent_level < base_level:
            root.items.append(item)
            last_item_at_level[base_level] = item
            continue

        parent_item = last_item_at_level[parent_level]

        if not parent_item.sub_lists or parent_item.sub_lists[-1].list_type != list_type:
            parent_item.sub_lists.append(ListBlock(list_type=list_type, items=[], slide_number=slide_number))

        parent_item.sub_lists[-1].items.append(item)
        last_item_at_level[level] = item

        for old_level in list(last_item_at_level):
            if old_level > level: del last_item_at_level[old_level]

    return root


# ============================================================
# TABLE
# ============================================================

def extract_table(shape, slide_number):
    rows = []

    for row in shape.table.rows:
        row_data = []

        for cell in row.cells:
            parts = [get_paragraph_text(p) for p in cell.text_frame.paragraphs if get_paragraph_text(p)]
            row_data.append("\n".join(parts))

        rows.append(row_data)

    if not rows or not any(cell.strip() for row in rows for cell in row): return None

    first_row = rows[0]
    non_empty = sum(1 for cell in first_row if cell.strip())
    bold = 0

    for index, cell in enumerate(shape.table.rows[0].cells):
        if index >= len(first_row) or not first_row[index].strip(): continue
        if any(run.font.bold is True for p in cell.text_frame.paragraphs for run in p.runs): bold += 1

    if non_empty and bold / non_empty >= 0.5: return TableBlock(headers=first_row, rows=rows[1:], slide_number=slide_number)

    return TableBlock(headers=[], rows=rows, slide_number=slide_number)


# ============================================================
# TITLE
# ============================================================

def get_slide_title_shape(slide):
    for shape in slide.shapes:
        if not shape.is_placeholder: continue

        try:
            if shape.placeholder_format.type in (PP_PLACEHOLDER.TITLE, PP_PLACEHOLDER.CENTER_TITLE): return shape
        except Exception:
            pass

    return None


# ============================================================
# PPTX -> NORMALIZED DOCUMENT
# ============================================================

def convert_pptx_to_normalized_document(file_path):
    prs = pptx.Presentation(file_path)
    path_lower = file_path.lower()

    if "aventro" in path_lower: enterprise = "Aventro Motors"
    elif "cendara" in path_lower: enterprise = "Cendara University"
    elif "cloudway" in path_lower: enterprise = "CloudWay-24"
    elif "velvera" in path_lower: enterprise = "Velvera Technologies"
    elif "zx bank" in path_lower or "zx_bank" in path_lower or "zx-bank" in path_lower or "zx" in path_lower: enterprise = "ZX Bank"
    else: enterprise = "Unknown"

    identity = enterprise + "::" + os.path.basename(file_path)
    document_id = hashlib.sha256(identity.encode("utf-8")).hexdigest()

    sections = []
    document_title = None
    manual_number_re = re.compile(r"^(\d+)\.\s+(.+)$")

    for slide_index, slide in enumerate(prs.slides):
        slide_number = slide_index + 1
        title_shape = get_slide_title_shape(slide)
        slide_title = get_shape_text(title_shape) if title_shape is not None else ""

        if document_title is None and slide_title: document_title = slide_title

        blocks = []
        shapes = sorted(slide.shapes, key=lambda shape: (int(shape.top), int(shape.left)))

        for shape in shapes:
            if title_shape is not None and shape._element is title_shape._element: continue

            # ------------------------------------------------
            # TABLE
            # ------------------------------------------------

            if getattr(shape, "has_table", False):
                table_block = extract_table(shape, slide_number)
                if table_block is not None: blocks.append(table_block)
                continue

            # ------------------------------------------------
            # TEXT
            # ------------------------------------------------

            if not shape.has_text_frame: continue

            buffered_list = []

            def flush_list():
                if not buffered_list: return
                list_block = build_list_block(buffered_list, slide_number)
                if list_block is not None and list_block.items: blocks.append(list_block)
                buffered_list.clear()

            for paragraph in shape.text_frame.paragraphs:
                text = get_paragraph_text(paragraph)
                if not text: continue

                is_list, level, list_type = get_list_info(paragraph, shape, slide)

                if is_list:
                    if buffered_list:
                        base_level = buffered_list[0][1]
                        previous_type = buffered_list[-1][2]
                        if level <= base_level and previous_type != list_type: flush_list()

                    buffered_list.append((text, level, list_type))
                    continue

                manual_match = manual_number_re.match(text)

                if manual_match:
                    if buffered_list and buffered_list[-1][2] != "ordered": flush_list()
                    buffered_list.append((manual_match.group(2).strip(), 0, "ordered"))
                    continue

                flush_list()
                blocks.append(TextBlock(text=text, slide_number=slide_number))

            flush_list()

        if slide_title or blocks: sections.append(Section(hierarchy=[slide_title] if slide_title else [], blocks=blocks))

    if document_title is None: document_title = os.path.splitext(os.path.basename(file_path))[0]

    metadata = DocumentMetadata(
        enterprise=enterprise,
        source_filename=os.path.basename(file_path),
        source_format="pptx",
        document_version=None,
        author=None,
        date=None,
        access_groups=[],
    )

    if not sections: sections = [Section(hierarchy=[], blocks=[])]

    return NormalizedDocument(document_id=document_id, title=document_title, metadata=metadata, sections=sections)