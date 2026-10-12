import os
import glob
from collections import Counter, defaultdict

import pptx
from pptx.enum.shapes import MSO_SHAPE_TYPE, PP_PLACEHOLDER
from pptx.oxml.ns import qn


# ============================================================
# BASIC HELPERS
# ============================================================

def add_example(examples, category, text, limit=5):
    if len(examples[category]) < limit:
        examples[category].append(text)


def get_text_preview(text, length=70):
    text = " ".join(text.split())
    return text[:length] + ("..." if len(text) > length else "")


# ============================================================
# BULLET / NUMBERING RESOLUTION
# ============================================================

def bullet_type_from_ppr(ppr):
    """
    Inspect a:pPr / a:lvlXpPr and return:

        bullet
        numbered
        none
        None -> not explicitly specified
    """

    if ppr is None:
        return None

    if ppr.find(qn("a:buNone")) is not None:
        return "none"

    if ppr.find(qn("a:buAutoNum")) is not None:
        return "numbered"

    if ppr.find(qn("a:buChar")) is not None:
        return "bullet"

    if ppr.find(qn("a:buBlip")) is not None:
        return "bullet"

    return None


def get_shape_level_style(shape, level):
    """Check list style defined directly on the current shape."""
    try:
        nodes = shape._element.xpath(f".//a:lstStyle/a:lvl{level + 1}pPr")
        return nodes[0] if nodes else None
    except Exception:
        return None


def get_layout_placeholder(slide, shape):
    """Find corresponding placeholder on slide layout using placeholder idx."""
    if not shape.is_placeholder:
        return None

    try:
        target_idx = shape.placeholder_format.idx
    except Exception:
        return None

    try:
        for placeholder in slide.slide_layout.placeholders:
            if placeholder.placeholder_format.idx == target_idx:
                return placeholder
    except Exception:
        pass

    return None


def get_master_level_style(slide, shape, level):
    """
    PowerPoint bullets can be inherited from the slide master.

    Decide which master text style to inspect based on placeholder type.
    """

    try:
        master = slide.slide_layout.slide_master
    except Exception:
        return None

    style_name = "otherStyle"

    if shape.is_placeholder:
        try:
            ph_type = shape.placeholder_format.type

            if ph_type in (PP_PLACEHOLDER.TITLE, PP_PLACEHOLDER.CENTER_TITLE):
                style_name = "titleStyle"

            elif ph_type in (
                PP_PLACEHOLDER.BODY,
                PP_PLACEHOLDER.OBJECT,
                PP_PLACEHOLDER.SUBTITLE,
            ):
                style_name = "bodyStyle"

        except Exception:
            pass

    try:
        nodes = master.element.xpath(
            f".//p:txStyles/p:{style_name}/a:lvl{level + 1}pPr"
        )

        return nodes[0] if nodes else None

    except Exception:
        return None


def get_effective_bullet_type(paragraph, shape, slide):
    """
    Resolve PowerPoint bullet/numbering using priority:

    paragraph
        ->
    shape list style
        ->
    layout placeholder
        ->
    slide master
    """

    level = paragraph.level or 0

    # 1. Paragraph-local properties
    direct = bullet_type_from_ppr(paragraph._p.pPr)

    if direct is not None:
        return direct

    # 2. Current shape list style
    shape_style = get_shape_level_style(shape, level)
    result = bullet_type_from_ppr(shape_style)

    if result is not None:
        return result

    # 3. Corresponding layout placeholder
    layout_placeholder = get_layout_placeholder(slide, shape)

    if layout_placeholder is not None:
        layout_style = get_shape_level_style(layout_placeholder, level)
        result = bullet_type_from_ppr(layout_style)

        if result is not None:
            return result

    # 4. Slide master
    master_style = get_master_level_style(slide, shape, level)
    result = bullet_type_from_ppr(master_style)

    if result is not None:
        return result

    # Nested paragraph is still suspicious/list-like even if
    # inherited bullet definition cannot be resolved.
    if level > 0:
        return "list-like"

    return None


# ============================================================
# HYPERLINK INSPECTION
# ============================================================

def inspect_text_frame_hyperlinks(text_frame, doc_id, slide_num, doc_flags, slide_flags, examples):
    count = 0

    for paragraph in text_frame.paragraphs:
        for run in paragraph.runs:
            try:
                address = run.hyperlink.address
            except Exception:
                address = None

            if address:
                count += 1
                doc_flags[doc_id].add("hyperlinks")
                slide_flags[(doc_id, slide_num)].add("hyperlinks")

                add_example(
                    examples,
                    "HYPERLINK",
                    f"{doc_id} Slide {slide_num}: '{run.text}' -> {address}",
                )

    return count


# ============================================================
# MAIN INSPECTION
# ============================================================

def inspect_pptx_corpus():
    print("=" * 72)
    print("PPTX CORPUS STRUCTURAL INSPECTION")
    print("=" * 72)

    corpus_dir = os.path.abspath(
        "Dataset/RAG-Multi-Corpus-main/RAG-Multi-Corpus-main/datasets"
    )

    if not os.path.exists(corpus_dir):
        corpus_dir = os.path.abspath(
            "Dataset/RAG-Multi-Corpus-main/RAG-Multi-Corpus-main/datasets"
        )

    pptx_files = glob.glob(
        os.path.join(corpus_dir, "**", "*.pptx"),
        recursive=True,
    )

    stats = Counter()
    doc_flags = defaultdict(set)
    slide_flags = defaultdict(set)

    table_stats = Counter()
    list_stats = Counter()
    shape_stats = Counter()

    slide_counts = []
    examples = defaultdict(list)

    stats["Total PPTX files"] = len(pptx_files)

    # ========================================================
    # FILE LOOP
    # ========================================================

    for file_path in pptx_files:
        doc_id = os.path.basename(file_path)

        try:
            prs = pptx.Presentation(file_path)
            stats["Readable"] += 1

        except Exception as e:
            stats["Failed"] += 1
            add_example(examples, "FAILED", f"{doc_id}: {e}")
            continue

        slide_counts.append(len(prs.slides))

        # ====================================================
        # SLIDE LOOP
        # ====================================================

        for slide_index, slide in enumerate(prs.slides):
            slide_num = slide_index + 1
            slide_key = (doc_id, slide_num)

            title_found = False
            body_found = False
            textbox_count = 0
            top_level_shape_summaries = []

            # =================================================
            # SHAPE LOOP
            # =================================================

            for shape_index, shape in enumerate(slide.shapes):
                shape_stats["total shapes"] += 1

                x = int(shape.left)
                y = int(shape.top)
                width = int(shape.width)
                height = int(shape.height)

                shape_type = str(shape.shape_type)
                text_preview = ""

                # ---------------------------------------------
                # PLACEHOLDER TYPE
                # ---------------------------------------------

                if shape.is_placeholder:
                    try:
                        ph_type = shape.placeholder_format.type
                    except Exception:
                        ph_type = None

                    if ph_type in (PP_PLACEHOLDER.TITLE, PP_PLACEHOLDER.CENTER_TITLE):
                        title_found = True
                        doc_flags[doc_id].add("title placeholders")
                        slide_flags[slide_key].add("title placeholders")

                    elif ph_type == PP_PLACEHOLDER.SUBTITLE:
                        doc_flags[doc_id].add("subtitle placeholders")
                        slide_flags[slide_key].add("subtitle placeholders")

                    elif ph_type == PP_PLACEHOLDER.BODY:
                        body_found = True
                        doc_flags[doc_id].add("body placeholders")
                        slide_flags[slide_key].add("body placeholders")

                    elif ph_type == PP_PLACEHOLDER.OBJECT:
                        doc_flags[doc_id].add("content/object placeholders")
                        slide_flags[slide_key].add("content/object placeholders")

                # ---------------------------------------------
                # TABLE
                #
                # IMPORTANT:
                # use has_table rather than shape_type.
                # ---------------------------------------------

                if getattr(shape, "has_table", False):
                    doc_flags[doc_id].add("tables")
                    slide_flags[slide_key].add("tables")

                    table_stats["total tables"] += 1

                    rows = len(shape.table.rows)
                    cols = len(shape.table.columns)

                    if cols == 1:
                        table_stats["1-column tables"] += 1
                    elif cols == 2:
                        table_stats["2-column tables"] += 1
                    else:
                        table_stats["3+ column tables"] += 1

                    empty_cells = 0

                    for row in shape.table.rows:
                        for cell in row.cells:
                            if not cell.text.strip():
                                empty_cells += 1

                            inspect_text_frame_hyperlinks(
                                cell.text_frame,
                                doc_id,
                                slide_num,
                                doc_flags,
                                slide_flags,
                                examples,
                            )

                    if empty_cells:
                        table_stats["tables with empty cells"] += 1
                        table_stats["total empty cells"] += empty_cells

                        doc_flags[doc_id].add("tables with empty cells")
                        slide_flags[slide_key].add("tables with empty cells")

                    preview_rows = []
                    for i, row in enumerate(shape.table.rows):
                        if i >= 3: break
                        preview_rows.append(" | ".join(get_text_preview(cell.text, 40) for cell in row.cells))
                    add_example(
                        examples,
                        "TABLE",
                        f"{doc_id} Slide {slide_num}: {rows}x{cols} -> {' / '.join(preview_rows)}",
                    )

                # ---------------------------------------------
                # CHART
                # ---------------------------------------------

                elif getattr(shape, "has_chart", False):
                    doc_flags[doc_id].add("charts")
                    slide_flags[slide_key].add("charts")
                    shape_stats["charts"] += 1

                # ---------------------------------------------
                # IMAGE
                # ---------------------------------------------

                elif shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
                    doc_flags[doc_id].add("images/pictures")
                    slide_flags[slide_key].add("images/pictures")
                    shape_stats["pictures"] += 1

                    add_example(
                        examples,
                        "IMAGE",
                        f"{doc_id} Slide {slide_num}: x={x}, y={y}, w={width}, h={height}",
                    )

                # ---------------------------------------------
                # GROUP
                # ---------------------------------------------

                elif shape.shape_type == MSO_SHAPE_TYPE.GROUP:
                    doc_flags[doc_id].add("grouped shapes")
                    slide_flags[slide_key].add("grouped shapes")
                    shape_stats["grouped shapes"] += 1

                    add_example(
                        examples,
                        "GROUPED/DIAGRAM SHAPE",
                        f"{doc_id} Slide {slide_num}: grouped shape with {len(shape.shapes)} child shapes",
                    )

                # ---------------------------------------------
                # TEXT BOX
                # ---------------------------------------------

                if shape.shape_type == MSO_SHAPE_TYPE.TEXT_BOX:
                    textbox_count += 1
                    doc_flags[doc_id].add("normal text boxes")
                    slide_flags[slide_key].add("normal text boxes")

                # ---------------------------------------------
                # TEXT + LISTS
                # ---------------------------------------------

                if shape.has_text_frame:
                    full_text = shape.text.strip()

                    if full_text:
                        text_preview = get_text_preview(full_text)

                    non_empty_paragraphs = [
                        p for p in shape.text_frame.paragraphs if p.text.strip()
                    ]

                    max_level_in_shape = 0
                    resolved_list_paragraphs = 0

                    for paragraph in non_empty_paragraphs:
                        level = paragraph.level or 0
                        max_level_in_shape = max(max_level_in_shape, level)

                        bullet_type = get_effective_bullet_type(
                            paragraph,
                            shape,
                            slide,
                        )

                        if bullet_type == "bullet":
                            list_stats["bullet paragraphs"] += 1
                            list_stats[f"bullet level {level}"] += 1

                            doc_flags[doc_id].add("bullet paragraphs")
                            slide_flags[slide_key].add("bullet paragraphs")
                            resolved_list_paragraphs += 1

                        elif bullet_type == "numbered":
                            list_stats["numbered paragraphs"] += 1
                            list_stats[f"numbered level {level}"] += 1

                            doc_flags[doc_id].add("numbered paragraphs")
                            slide_flags[slide_key].add("numbered paragraphs")
                            resolved_list_paragraphs += 1

                        elif bullet_type == "list-like":
                            list_stats["unresolved nested list-like paragraphs"] += 1
                            list_stats[f"list-like level {level}"] += 1

                            doc_flags[doc_id].add("nested/list-like paragraphs")
                            slide_flags[slide_key].add("nested/list-like paragraphs")
                            resolved_list_paragraphs += 1

                        if level > 0:
                            list_stats["paragraphs level > 0"] += 1
                            list_stats[f"all paragraph level {level}"] += 1

                            doc_flags[doc_id].add("nested paragraph levels")
                            slide_flags[slide_key].add("nested paragraph levels")

                    if resolved_list_paragraphs:
                        list_stats["shapes containing lists"] += 1

                    if max_level_in_shape > 0:
                        add_example(
                            examples,
                            "NESTED BULLETS",
                            f"{doc_id} Slide {slide_num}: max paragraph level={max_level_in_shape}, text='{text_preview}'",
                        )

                    inspect_text_frame_hyperlinks(
                        shape.text_frame,
                        doc_id,
                        slide_num,
                        doc_flags,
                        slide_flags,
                        examples,
                    )

                top_level_shape_summaries.append(
                    {
                        "idx": shape_index,
                        "type": shape_type,
                        "x": x,
                        "y": y,
                        "w": width,
                        "h": height,
                        "text": text_preview,
                    }
                )

            # =================================================
            # SLIDE-LEVEL EXAMPLES
            # =================================================

            if title_found and body_found:
                add_example(
                    examples,
                    "TITLE + BODY",
                    f"{doc_id} Slide {slide_num}: title + body placeholders",
                )

            if textbox_count >= 3:
                sorted_shapes = sorted(
                    top_level_shape_summaries,
                    key=lambda item: (item["y"], item["x"]),
                )

                lines = [f"{doc_id} Slide {slide_num}:"]

                for info in sorted_shapes:
                    lines.append(
                        f"  Shape {info['idx']}: type={info['type']}, "
                        f"x={info['x']}, y={info['y']}, "
                        f"w={info['w']}, h={info['h']} | {info['text']}"
                    )

                add_example(
                    examples,
                    "MULTI-TEXTBOX SLIDE",
                    "\n".join(lines),
                )

    # ========================================================
    # REPORT
    # ========================================================

    print("\n--- HIGH-LEVEL STATS ---")
    print(f"Total PPTX files: {stats['Total PPTX files']}")
    print(f"Readable: {stats['Readable']}")
    print(f"Failed: {stats['Failed']}")

    if slide_counts:
        total_slides = sum(slide_counts)

        print(f"Total slides: {total_slides}")
        print(f"Min slides: {min(slide_counts)}")
        print(f"Max slides: {max(slide_counts)}")
        print(f"Average slides: {total_slides / len(slide_counts):.1f}")

    # ========================================================
    # DOCUMENT FEATURES
    # ========================================================

    print("\n--- PRESENCE COUNTS (DOCUMENTS) ---")

    doc_feature_counts = Counter(
        feature
        for features in doc_flags.values()
        for feature in features
    )

    for feature, count in doc_feature_counts.most_common():
        print(f"{feature}: {count}")

    # ========================================================
    # SLIDE FEATURES
    # ========================================================

    print("\n--- PRESENCE COUNTS (SLIDES) ---")

    slide_feature_counts = Counter(
        feature
        for features in slide_flags.values()
        for feature in features
    )

    for feature, count in slide_feature_counts.most_common():
        print(f"{feature}: {count}")

    # ========================================================
    # TABLES
    # ========================================================

    print("\n--- TABLE STATS ---")

    if table_stats:
        for key, value in table_stats.items():
            print(f"{key}: {value}")
    else:
        print("No native PowerPoint tables detected.")

    # ========================================================
    # LISTS / BULLETS
    # ========================================================

    print("\n--- LIST / BULLET STATS ---")

    if list_stats:
        for key, value in list_stats.items():
            print(f"{key}: {value}")
    else:
        print("No bullet/numbered/list-level structure detected.")

    # ========================================================
    # OTHER SHAPES
    # ========================================================

    print("\n--- OTHER SHAPE STATS ---")

    if shape_stats:
        for key, value in shape_stats.items():
            print(f"{key}: {value}")

    # ========================================================
    # EXAMPLES
    # ========================================================

    print("\n--- REPRESENTATIVE SLIDE STRUCTURES ---")

    preferred = [
        "TITLE + BODY",
        "NESTED BULLETS",
        "TABLE",
        "HYPERLINK",
        "IMAGE",
        "MULTI-TEXTBOX SLIDE",
        "GROUPED/DIAGRAM SHAPE",
        "FAILED",
    ]

    for category in preferred:
        values = examples.get(category, [])

        if not values:
            continue

        print(f"\n[{category}]")

        for value in values:
            print(f"  - {value}")

    print("\n" + "=" * 72)
    print("PPTX INSPECTION COMPLETE")
    print("=" * 72)


if __name__ == "__main__":
    inspect_pptx_corpus()