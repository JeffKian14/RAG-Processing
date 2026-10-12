import os
import glob
import re
from collections import Counter, defaultdict

import docx
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn


# ============================================================
# CONFIG
# ============================================================

MAX_EXAMPLES = 5

MANUAL_BULLET_RE = re.compile(
    r"^([•\-–—*✓➢►▪◦])\s+"
)

MANUAL_NUMBER_RE = re.compile(
    r"^("
    r"\d+\."
    r"|\d+\)"
    r"|[a-zA-Z]\."
    r"|[a-zA-Z]\)"
    r"|\([a-zA-Z0-9]+\)"
    r")\s+"
)

MULTI_LEVEL_NUMBER_RE = re.compile(
    r"^\d+\.\d+(?:\.\d+)*[\.\)]?\s+"
)


# ============================================================
# HELPERS
# ============================================================

def add_example(examples, category, text):
    text = str(text).strip()

    if not text:
        return

    if len(examples[category]) >= MAX_EXAMPLES:
        return

    text = re.sub(r"\s+", " ", text)

    examples[category].append(
        text[:250]
    )


def safe_style_name(paragraph):
    try:
        if paragraph.style:
            return paragraph.style.name or ""
    except Exception:
        pass

    return ""


def paragraph_font_sizes(paragraph):
    sizes = []

    for run in paragraph.runs:
        try:
            if run.font.size:
                sizes.append(
                    float(run.font.size.pt)
                )
        except Exception:
            pass

    return sizes


def paragraph_max_font_size(paragraph):
    sizes = paragraph_font_sizes(
        paragraph
    )

    return max(sizes) if sizes else None


def paragraph_is_bold(paragraph):
    return any(
        run.bold is True
        for run in paragraph.runs
    )


def paragraph_has_colour(paragraph):
    for run in paragraph.runs:

        try:
            color = run.font.color

            if (
                color is not None
                and color.rgb is not None
            ):
                return True

        except Exception:
            pass

    return False


def paragraph_is_centered(paragraph):
    return (
        paragraph.alignment
        == WD_ALIGN_PARAGRAPH.CENTER
    )


# ============================================================
# NATIVE WORD NUMBERING
# ============================================================

def get_num_properties(paragraph):
    """
    Returns:

        num_id
        ilvl

    for a native Word-numbered paragraph.

    Returns (None, None) if paragraph does not
    contain numPr.
    """

    p_pr = paragraph._p.pPr

    if p_pr is None:
        return None, None

    num_pr = p_pr.numPr

    if num_pr is None:
        return None, None

    num_id = None
    ilvl = 0

    if num_pr.numId is not None:
        try:
            num_id = int(
                num_pr.numId.val
            )
        except Exception:
            pass

    if num_pr.ilvl is not None:
        try:
            ilvl = int(
                num_pr.ilvl.val
            )
        except Exception:
            pass

    return num_id, ilvl


def build_numbering_map(doc):
    """
    Resolve:

        numId
            ->
        abstractNumId
            ->
        ilvl
            ->
        numFmt

    Example numFmt values:

        bullet
        decimal
        lowerLetter
        upperLetter
        lowerRoman
        upperRoman
    """

    result = {}

    try:
        numbering_part = (
            doc.part.numbering_part
        )

        root = numbering_part.element

    except Exception:
        return result

    # --------------------------------------------
    # abstractNum definitions
    # --------------------------------------------

    abstract_levels = {}

    for abstract_num in root.findall(
        qn("w:abstractNum")
    ):

        abstract_id = abstract_num.get(
            qn("w:abstractNumId")
        )

        if abstract_id is None:
            continue

        level_formats = {}

        for lvl in abstract_num.findall(
            qn("w:lvl")
        ):

            ilvl = lvl.get(
                qn("w:ilvl")
            )

            if ilvl is None:
                continue

            num_fmt = lvl.find(
                qn("w:numFmt")
            )

            if num_fmt is not None:

                fmt = num_fmt.get(
                    qn("w:val")
                )

            else:

                fmt = "unknown"

            try:
                level_formats[
                    int(ilvl)
                ] = fmt

            except ValueError:
                continue

        abstract_levels[
            str(abstract_id)
        ] = level_formats

    # --------------------------------------------
    # numId -> abstractNumId
    # --------------------------------------------

    for num in root.findall(
        qn("w:num")
    ):

        num_id = num.get(
            qn("w:numId")
        )

        abstract_ref = num.find(
            qn("w:abstractNumId")
        )

        if (
            num_id is None
            or abstract_ref is None
        ):
            continue

        abstract_id = abstract_ref.get(
            qn("w:val")
        )

        if abstract_id is None:
            continue

        try:
            numeric_num_id = int(num_id)
        except ValueError:
            continue

        result[
            numeric_num_id
        ] = abstract_levels.get(
            str(abstract_id),
            {}
        )

    return result


def get_native_list_format(
    paragraph,
    numbering_map,
):
    """
    Returns:

        (format, level)

    Example:

        ("bullet", 0)
        ("decimal", 1)

    Returns:

        (None, None)

    if not a native numbered paragraph.
    """

    num_id, ilvl = get_num_properties(
        paragraph
    )

    if num_id is None:
        return None, None

    levels = numbering_map.get(
        num_id,
        {}
    )

    fmt = levels.get(
        ilvl,
        "unknown"
    )

    return fmt, ilvl


# ============================================================
# FRONT MATTER DETECTION
# ============================================================

def detect_front_matter(doc):
    """
    Conservative detector.

    Strong front-matter clues:

        Title style
        centered
        font > 14 pt

    Bold / colour alone are NOT enough.

    Only considers the beginning of the document.
    """

    meaningful = [
        p
        for p in doc.paragraphs[:12]
        if p.text.strip()
    ]

    if not meaningful:
        return []

    candidates = []

    for index, paragraph in enumerate(
        meaningful[:6]
    ):

        text = paragraph.text.strip()

        style_name = safe_style_name(
            paragraph
        )

        max_size = paragraph_max_font_size(
            paragraph
        )

        is_title = (
            style_name.lower()
            == "title"
        )

        is_centered = (
            paragraph_is_centered(
                paragraph
            )
        )

        is_large = (
            max_size is not None
            and max_size > 14.0
        )

        strong_special = (
            is_title
            or is_centered
            or is_large
        )

        # First few lines only.
        if strong_special:

            candidates.append(text)

        else:

            # Once normal body content starts,
            # stop treating later lines as front matter.
            if candidates:
                break

            # Allow first normal line before a strong title
            # only minimally; otherwise stop quickly.
            if index >= 2:
                break

    # Need >= 2 lines because a single heading does
    # not constitute a title block.
    if len(candidates) >= 2:
        return candidates

    return []


# ============================================================
# TABLE HELPERS
# ============================================================

def table_has_merged_cells(table):
    """
    Detect horizontal gridSpan or vertical merge.
    """

    for row in table.rows:

        for cell in row.cells:

            tc_pr = cell._tc.tcPr

            if tc_pr is None:
                continue

            if (
                tc_pr.gridSpan is not None
                or tc_pr.vMerge is not None
            ):
                return True

    return False


def table_empty_cell_count(table):
    count = 0

    for row in table.rows:

        for cell in row.cells:

            if not cell.text.strip():
                count += 1

    return count


def table_is_completely_empty(table):
    return all(
        not cell.text.strip()
        for row in table.rows
        for cell in row.cells
    )


def table_text_preview(table):
    values = []

    for row in table.rows[:3]:

        row_values = [
            re.sub(
                r"\s+",
                " ",
                cell.text.strip(),
            )
            for cell in row.cells
        ]

        values.append(
            " | ".join(row_values)
        )

    return " / ".join(values)


# ============================================================
# XML FEATURE DETECTION
# ============================================================

def detect_xml_features(doc):
    flags = set()

    element = doc.element

    try:
        if element.xpath("//w:ins"):
            flags.add(
                "tracked changes"
            )

        if element.xpath("//w:del"):
            flags.add(
                "tracked changes"
            )

    except Exception:
        pass

    try:
        if element.xpath(
            "//w:commentReference"
        ):
            flags.add(
                "comments"
            )

    except Exception:
        pass

    try:
        if element.xpath(
            "//w:txbxContent"
        ):
            flags.add(
                "text boxes"
            )

    except Exception:
        pass

    try:
        if element.xpath(
            "//w:footnoteReference"
        ):
            flags.add(
                "footnotes"
            )

        if element.xpath(
            "//w:endnoteReference"
        ):
            flags.add(
                "endnotes"
            )

    except Exception:
        pass

    try:
        if element.xpath(
            "//w:hyperlink"
        ):
            flags.add(
                "hyperlinks"
            )

    except Exception:
        pass

    try:
        if element.xpath(
            "//w:drawing"
        ):
            flags.add(
                "drawings/images"
            )

    except Exception:
        pass

    try:
        instr_texts = element.xpath(
            "//w:instrText"
        )

        for instr in instr_texts:

            text = (
                instr.text or ""
            ).upper()

            if "TOC" in text:
                flags.add(
                    "table of contents"
                )
                break

    except Exception:
        pass

    return flags


# ============================================================
# MAIN INSPECTION
# ============================================================

def inspect_docx_corpus():

    print(
        "=" * 72
    )

    print(
        "DOCX CORPUS STRUCTURAL INSPECTION"
    )

    print(
        "=" * 72
    )

    # --------------------------------------------------------
    # Dataset path
    # --------------------------------------------------------

    corpus_dir = os.path.abspath(
        "Dataset/"
        "RAG-Multi-Corpus-main/"
        "RAG-Multi-Corpus-main/"
        "datasets"
    )

    if not os.path.exists(
        corpus_dir
    ):

        corpus_dir = os.path.abspath(
            "Dataset/"
            "RAG-Multi-Corpus-main/"
            "RAG-Multi_Corpus-main/"
            "datasets"
        )

    docx_files = glob.glob(
        os.path.join(
            corpus_dir,
            "**",
            "*.docx",
        ),
        recursive=True,
    )

    stats = Counter()

    style_counts = Counter()

    numbering_formats = Counter()

    numbering_documents = Counter()

    max_list_level = 0

    examples = defaultdict(list)

    stats[
        "Total DOCX files"
    ] = len(docx_files)

    # Table global statistics
    total_tables = 0
    total_empty_cells = 0

    # --------------------------------------------------------
    # Inspect files
    # --------------------------------------------------------

    for file_path in docx_files:

        filename = os.path.basename(
            file_path
        )

        try:

            doc = docx.Document(
                file_path
            )

            stats[
                "Readable"
            ] += 1

        except Exception as exc:

            stats[
                "Failed"
            ] += 1

            add_example(
                examples,
                "FAILED",
                f"{filename}: {exc}",
            )

            continue

        document_flags = set()

        # ====================================================
        # XML FEATURES
        # ====================================================

        xml_flags = detect_xml_features(
            doc
        )

        for flag in xml_flags:
            document_flags.add(
                flag
            )

        # ====================================================
        # FRONT MATTER
        # ====================================================

        front_matter = (
            detect_front_matter(
                doc
            )
        )

        if front_matter:

            document_flags.add(
                "strong front matter"
            )

            add_example(
                examples,
                "FRONT MATTER",
                (
                    filename
                    + " -> "
                    + " | ".join(
                        front_matter
                    )
                ),
            )

        # ====================================================
        # NUMBERING MAP
        # ====================================================

        numbering_map = (
            build_numbering_map(
                doc
            )
        )

        formats_seen_in_doc = set()

        # ====================================================
        # PARAGRAPHS
        # ====================================================

        if doc.paragraphs:

            document_flags.add(
                "paragraphs"
            )

        for paragraph in doc.paragraphs:

            text = (
                paragraph.text.strip()
            )

            if not text:
                continue

            style_name = (
                safe_style_name(
                    paragraph
                )
                or "Unknown"
            )

            style_counts[
                style_name
            ] += 1

            # ----------------------------------------------
            # Semantic styles
            # ----------------------------------------------

            style_lower = (
                style_name.lower()
            )

            if style_lower == "title":

                document_flags.add(
                    "Title style"
                )

            elif style_lower == "heading 1":

                document_flags.add(
                    "Heading 1"
                )

            elif style_lower == "heading 2":

                document_flags.add(
                    "Heading 2"
                )

            elif style_lower == "heading 3":

                document_flags.add(
                    "Heading 3"
                )

            elif style_lower.startswith(
                "heading"
            ):

                document_flags.add(
                    "Heading 4+"
                )

            # ----------------------------------------------
            # Formatting evidence
            # ----------------------------------------------

            if paragraph_is_bold(
                paragraph
            ):

                document_flags.add(
                    "bold paragraphs"
                )

            if paragraph_has_colour(
                paragraph
            ):

                document_flags.add(
                    "colored paragraphs"
                )

            max_font = paragraph_max_font_size(
                paragraph
            )

            if (
                max_font is not None
                and max_font > 14
            ):

                document_flags.add(
                    "large-font paragraphs"
                )

            if paragraph_is_centered(
                paragraph
            ):

                document_flags.add(
                    "centered paragraphs"
                )

            # ----------------------------------------------
            # Native Word lists
            # ----------------------------------------------

            list_fmt, level = (
                get_native_list_format(
                    paragraph,
                    numbering_map,
                )
            )

            if list_fmt is not None:

                document_flags.add(
                    "native lists"
                )

                numbering_formats[
                    list_fmt
                ] += 1

                formats_seen_in_doc.add(
                    list_fmt
                )

                max_list_level = max(
                    max_list_level,
                    level or 0,
                )

                add_example(
                    examples,
                    f"NATIVE {list_fmt.upper()}",
                    (
                        f"{filename} "
                        f"[level {level}] "
                        f"{text}"
                    ),
                )

            # ----------------------------------------------
            # Manual numbering / bullets
            # ----------------------------------------------

            if MULTI_LEVEL_NUMBER_RE.match(
                text
            ):

                document_flags.add(
                    "manual multi-level numbering"
                )

                add_example(
                    examples,
                    "MULTI-LEVEL NUMBERING",
                    f"{filename} -> {text}",
                )

            elif MANUAL_NUMBER_RE.match(
                text
            ):

                document_flags.add(
                    "manual numbering"
                )

                add_example(
                    examples,
                    "MANUAL NUMBERING",
                    f"{filename} -> {text}",
                )

            elif MANUAL_BULLET_RE.match(
                text
            ):

                document_flags.add(
                    "manual bullets"
                )

                add_example(
                    examples,
                    "MANUAL BULLET",
                    f"{filename} -> {text}",
                )

            # ----------------------------------------------
            # Captions
            # ----------------------------------------------

            if (
                style_lower == "caption"
                or text.lower().startswith(
                    "figure "
                )
                or text.lower().startswith(
                    "table "
                )
            ):

                document_flags.add(
                    "captions"
                )

        # Count how many documents contain each
        # native numbering format.
        for fmt in formats_seen_in_doc:

            numbering_documents[
                fmt
            ] += 1

        # ====================================================
        # TABLES
        # ====================================================

        if doc.tables:

            document_flags.add(
                "tables"
            )

        for table in doc.tables:

            total_tables += 1

            rows = len(
                table.rows
            )

            cols = len(
                table.columns
            )

            if cols == 1:

                stats[
                    "1-column tables"
                ] += 1

            elif cols == 2:

                stats[
                    "2-column tables"
                ] += 1

            elif cols >= 3:

                stats[
                    "3+-column tables"
                ] += 1

            if (
                rows == 1
                and cols == 1
            ):

                stats[
                    "One-cell tables"
                ] += 1

                document_flags.add(
                    "one-cell tables"
                )

                add_example(
                    examples,
                    "ONE-CELL TABLE",
                    (
                        filename
                        + " -> "
                        + table_text_preview(
                            table
                        )
                    ),
                )

            if table_has_merged_cells(
                table
            ):

                stats[
                    "Tables with merged cells"
                ] += 1

                document_flags.add(
                    "merged tables"
                )

                add_example(
                    examples,
                    "MERGED TABLE",
                    (
                        filename
                        + " -> "
                        + table_text_preview(
                            table
                        )
                    ),
                )

            empty_cells = (
                table_empty_cell_count(
                    table
                )
            )

            total_empty_cells += (
                empty_cells
            )

            if empty_cells > 0:

                stats[
                    "Tables containing empty cells"
                ] += 1

                document_flags.add(
                    "tables with empty cells"
                )

            if table_is_completely_empty(
                table
            ):

                stats[
                    "Completely empty tables"
                ] += 1

            add_example(
                examples,
                "TABLE",
                (
                    filename
                    + " -> "
                    + table_text_preview(
                        table
                    )
                ),
            )

        # ====================================================
        # DOCUMENT FLAGS
        # ====================================================

        for flag in document_flags:

            stats[
                f"Documents containing {flag}"
            ] += 1

    # ========================================================
    # OUTPUT
    # ========================================================

    print()
    print(
        "--- HIGH-LEVEL ---"
    )

    print(
        f"Total DOCX files: "
        f"{stats['Total DOCX files']}"
    )

    print(
        f"Readable: "
        f"{stats['Readable']}"
    )

    print(
        f"Failed: "
        f"{stats['Failed']}"
    )

    # ========================================================
    # DOCUMENT STRUCTURE
    # ========================================================

    print()
    print(
        "--- DOCUMENT STRUCTURE ---"
    )

    feature_order = [
        "paragraphs",
        "Title style",
        "Heading 1",
        "Heading 2",
        "Heading 3",
        "Heading 4+",
        "strong front matter",
        "native lists",
        "manual numbering",
        "manual bullets",
        "manual multi-level numbering",
        "tables",
        "merged tables",
        "tables with empty cells",
        "one-cell tables",
        "hyperlinks",
        "drawings/images",
        "captions",
        "text boxes",
        "table of contents",
        "tracked changes",
        "comments",
        "footnotes",
        "endnotes",
    ]

    for feature in feature_order:

        count = stats.get(
            f"Documents containing {feature}",
            0,
        )

        print(
            f"{feature}: {count} documents"
        )

    # ========================================================
    # NUMBERING
    # ========================================================

    print()
    print(
        "--- NATIVE WORD LIST FORMATS ---"
    )

    if numbering_formats:

        for fmt, count in (
            numbering_formats.most_common()
        ):

            docs = numbering_documents[
                fmt
            ]

            print(
                f"{fmt}: "
                f"{count} paragraphs "
                f"across {docs} documents"
            )

    else:

        print(
            "No native numbering resolved."
        )

    print(
        f"Maximum native nesting level "
        f"(ilvl): {max_list_level}"
    )

    # ========================================================
    # TABLES
    # ========================================================

    print()
    print(
        "--- TABLE STRUCTURE ---"
    )

    print(
        f"Total tables: "
        f"{total_tables}"
    )

    print(
        f"1-column tables: "
        f"{stats['1-column tables']}"
    )

    print(
        f"2-column tables: "
        f"{stats['2-column tables']}"
    )

    print(
        f"3+ column tables: "
        f"{stats['3+-column tables']}"
    )

    print(
        f"One-cell tables: "
        f"{stats['One-cell tables']}"
    )

    print(
        f"Tables with merged cells: "
        f"{stats['Tables with merged cells']}"
    )

    print(
        f"Tables containing empty cells: "
        f"{stats['Tables containing empty cells']}"
    )

    print(
        f"Completely empty tables: "
        f"{stats['Completely empty tables']}"
    )

    print(
        f"Total empty table cells: "
        f"{total_empty_cells}"
    )

    # ========================================================
    # STYLES
    # ========================================================

    print()
    print(
        "--- TOP 20 PARAGRAPH STYLES ---"
    )

    for style_name, count in (
        style_counts.most_common(20)
    ):

        print(
            f"{style_name}: "
            f"{count}"
        )

    # ========================================================
    # EXAMPLES
    # ========================================================

    print()
    print(
        "--- REPRESENTATIVE EXAMPLES ---"
    )

    preferred_categories = [
        "FRONT MATTER",
        "NATIVE BULLET",
        "NATIVE DECIMAL",
        "NATIVE LOWERLETTER",
        "NATIVE UPPERLETTER",
        "NATIVE LOWERROMAN",
        "NATIVE UPPERROMAN",
        "NATIVE UNKNOWN",
        "MANUAL BULLET",
        "MANUAL NUMBERING",
        "MULTI-LEVEL NUMBERING",
        "TABLE",
        "MERGED TABLE",
        "ONE-CELL TABLE",
    ]

    printed = set()

    for category in preferred_categories:

        entries = examples.get(
            category,
            []
        )

        if not entries:
            continue

        printed.add(category)

        print()
        print(
            f"[{category}]"
        )

        for example in entries:

            print(
                f"  - {example}"
            )

    # Print remaining interesting categories.
    for category, entries in examples.items():

        if (
            category in printed
            or not entries
        ):
            continue

        print()
        print(
            f"[{category}]"
        )

        for example in entries:

            print(
                f"  - {example}"
            )

    print()
    print(
        "=" * 72
    )

    print(
        "INSPECTION COMPLETE"
    )

    print(
        "=" * 72
    )


# ============================================================
# ENTRY POINT
# ============================================================

if __name__ == "__main__":
    inspect_docx_corpus()