import os
import glob
import re
from collections import Counter

import docx
from docx.oxml.ns import qn

from src.parsers.docx_parser import (
    convert_docx_to_normalized_document,
    get_paragraph_text,
    get_list_info,
    build_numbering_map,
    iter_block_items,
)

from src.schemas import TextBlock, TableBlock, ListBlock


MANUAL_NUMBER_RE = re.compile(r"^(\d+)\.\s+(.+)$")
MULTI_LEVEL_RE = re.compile(r"^\d+\.\d+(?:\.\d+)*[\.\)]?\s+")


def normalize_text(text):
    return re.sub(r"\s+", " ", text).strip()


# ============================================================
# NORMALIZED LIST FLATTENING
# ============================================================

def flatten_list_items(list_block, depth=0):
    results = []

    for item in list_block.items:
        results.append((normalize_text(item.text), list_block.list_type, depth))

        for sub_list in item.sub_lists:
            results.extend(flatten_list_items(sub_list, depth + 1))

    return results


def collect_normalized_content(doc):
    texts = []
    lists = []
    tables = []
    hierarchies = []

    for section in doc.sections:
        hierarchies.append(section.hierarchy)

        for block in section.blocks:
            if isinstance(block, TextBlock):
                texts.append(block.text)

            elif isinstance(block, ListBlock):
                lists.extend(flatten_list_items(block))

            elif isinstance(block, TableBlock):
                tables.append(block)

    return {
        "texts": texts,
        "lists": lists,
        "tables": tables,
        "hierarchies": hierarchies,
    }


# ============================================================
# TABLE HELPERS
# ============================================================

def source_table_rows(table):
    rows = []

    for row in table.rows:
        row_data = []

        for cell in row.cells:
            parts = []

            for paragraph in cell.paragraphs:
                text = get_paragraph_text(paragraph)

                if text:
                    parts.append(text)

            row_data.append("\n".join(parts))

        rows.append(row_data)

    return rows


def normalized_table_rows(tables):
    rows = []

    for table in tables:
        if table.headers:
            rows.append(table.headers)

        rows.extend(table.rows)

    return rows


def count_empty_cells(rows):
    return sum(1 for row in rows for cell in row if cell == "")


# ============================================================
# SOURCE LIST GROUPS
# ============================================================

def collect_source_list_groups(source_doc, numbering_map):
    groups = []
    current = []

    def flush():
        nonlocal current

        if not current:
            return

        base_level = current[0][1]
        normalized_group = []

        for text, ilvl, list_type in current:
            normalized_group.append((normalize_text(text), list_type, max(0, ilvl - base_level)))

        groups.append(normalized_group)
        current = []

    for block in iter_block_items(source_doc):
        if isinstance(block, docx.table.Table):
            flush()
            continue

        text = get_paragraph_text(block)

        if not text:
            continue

        style_name = block.style.name if block.style else ""

        if style_name.startswith("Heading") or style_name == "Title":
            flush()
            continue

        is_list, ilvl, list_type = get_list_info(block, source_doc, numbering_map)

        if is_list:
            if current:
                base_level = current[0][1]

                if ilvl <= base_level and current[-1][2] != list_type:
                    flush()

            current.append((text, ilvl, list_type))
            continue

        manual_match = MANUAL_NUMBER_RE.match(text)

        if manual_match:
            if current and current[-1][2] != "ordered":
                flush()

            current.append((manual_match.group(2).strip(), 0, "ordered"))
            continue

        flush()

    flush()

    return groups


def collect_normalized_list_groups(normalized_doc):
    groups = []

    for section in normalized_doc.sections:
        for block in section.blocks:
            if isinstance(block, ListBlock):
                groups.append(flatten_list_items(block))

    return groups


# ============================================================
# SOURCE / OUTPUT ORDER
# ============================================================

def source_atomic_sequence(source_doc, numbering_map):
    sequence = []

    for block in iter_block_items(source_doc):

        if isinstance(block, docx.text.paragraph.Paragraph):
            text = get_paragraph_text(block)

            if not text:
                continue

            style_name = block.style.name if block.style else ""

            if style_name.startswith("Heading") or style_name == "Title":
                continue

            is_list, _, _ = get_list_info(block, source_doc, numbering_map)

            if is_list:
                sequence.append(("content", normalize_text(text)))
                continue

            manual_match = MANUAL_NUMBER_RE.match(text)

            if manual_match:
                sequence.append(("content", normalize_text(manual_match.group(2))))
                continue

            sequence.append(("content", normalize_text(text)))

        elif isinstance(block, docx.table.Table):
            for row in source_table_rows(block):
                for cell in row:
                    sequence.append(("table_cell", normalize_text(cell)))

    return sequence


def normalized_atomic_sequence(normalized_doc):
    sequence = []

    for section in normalized_doc.sections:
        for block in section.blocks:

            if isinstance(block, TextBlock):
                sequence.append(("content", normalize_text(block.text)))

            elif isinstance(block, ListBlock):
                for text, _, _ in flatten_list_items(block):
                    sequence.append(("content", normalize_text(text)))

            elif isinstance(block, TableBlock):
                rows = []

                if block.headers:
                    rows.append(block.headers)

                rows.extend(block.rows)

                for row in rows:
                    for cell in row:
                        sequence.append(("table_cell", normalize_text(cell)))

    return sequence


# ============================================================
# SINGLE DOCUMENT VALIDATION
# ============================================================

def validate_single_document(file_path):
    source_doc = docx.Document(file_path)
    numbering_map = build_numbering_map(source_doc)

    normalized_doc = convert_docx_to_normalized_document(file_path)
    normalized = collect_normalized_content(normalized_doc)

    issues = []

    source_bullet_items = []
    source_decimal_items = []
    source_manual_items = []
    source_multi_level = []
    source_hyperlinks = []

    # --------------------------------------------------------
    # Lists
    # --------------------------------------------------------

    for paragraph in source_doc.paragraphs:
        text = get_paragraph_text(paragraph)

        if not text:
            continue

        style_name = paragraph.style.name if paragraph.style else ""

        if style_name.startswith("Heading"):
            continue

        is_list, _, list_type = get_list_info(paragraph, source_doc, numbering_map)

        if is_list:
            if list_type == "unordered":
                source_bullet_items.append(normalize_text(text))
            else:
                source_decimal_items.append(normalize_text(text))

            continue

        if MULTI_LEVEL_RE.match(text):
            source_multi_level.append(normalize_text(text))
            continue

        manual_match = MANUAL_NUMBER_RE.match(text)

        if manual_match:
            source_manual_items.append(normalize_text(manual_match.group(2)))

    normalized_unordered = {text for text, list_type, _ in normalized["lists"] if list_type == "unordered"}
    normalized_ordered = {text for text, list_type, _ in normalized["lists"] if list_type == "ordered"}

    missing_bullets = [x for x in source_bullet_items if x not in normalized_unordered]
    missing_decimal = [x for x in source_decimal_items if x not in normalized_ordered]
    missing_manual = [x for x in source_manual_items if x not in normalized_ordered]

    if missing_bullets:
        issues.append(f"Missing native bullet items: {len(missing_bullets)}")

    if missing_decimal:
        issues.append(f"Missing native decimal items: {len(missing_decimal)}")

    if missing_manual:
        issues.append(f"Missing manual ordered items: {len(missing_manual)}")

    # --------------------------------------------------------
    # Direct list nesting validation
    # --------------------------------------------------------

    source_groups = collect_source_list_groups(source_doc, numbering_map)
    normalized_groups = collect_normalized_list_groups(normalized_doc)

    if source_groups != normalized_groups:
        issues.append(f"List nesting/type mismatch: source_groups={len(source_groups)}, normalized_groups={len(normalized_groups)}")

    # --------------------------------------------------------
    # Hyperlinks
    # --------------------------------------------------------

    for paragraph in source_doc.paragraphs:
        try:
            hyperlink_nodes = paragraph._element.xpath(".//w:hyperlink")
        except Exception:
            hyperlink_nodes = []

        for hyperlink in hyperlink_nodes:
            text = "".join(node.text for node in hyperlink.iter(qn("w:t")) if node.text)
            text = normalize_text(text)

            if text:
                source_hyperlinks.append(text)

    normalized_strings = []

    normalized_strings.extend(normalize_text(x) for x in normalized["texts"])
    normalized_strings.extend(text for text, _, _ in normalized["lists"])

    for hierarchy in normalized["hierarchies"]:
        normalized_strings.extend(normalize_text(x) for x in hierarchy)

    if normalized_doc.title:
        normalized_strings.append(normalize_text(normalized_doc.title))

    for table in normalized["tables"]:
        rows = ([table.headers] if table.headers else []) + table.rows

        for row in rows:
            normalized_strings.extend(normalize_text(cell) for cell in row)

    combined_normalized = " ".join(normalized_strings)

    missing_hyperlinks = [text for text in source_hyperlinks if text not in combined_normalized]

    if missing_hyperlinks:
        issues.append(f"Missing hyperlink visible text: {len(missing_hyperlinks)}")

    # --------------------------------------------------------
    # Tables
    # --------------------------------------------------------

    source_tables = source_doc.tables
    output_tables = normalized["tables"]

    if len(source_tables) != len(output_tables):
        issues.append(f"Table count mismatch: source={len(source_tables)}, normalized={len(output_tables)}")

    src_rows = []

    for table in source_tables:
        src_rows.extend(source_table_rows(table))

    out_rows = normalized_table_rows(output_tables)

    source_empty_cells = count_empty_cells(src_rows)
    normalized_empty_cells = count_empty_cells(out_rows)

    if source_empty_cells != normalized_empty_cells:
        issues.append(f"Empty-cell mismatch: source={source_empty_cells}, normalized={normalized_empty_cells}")

    source_table_values = [normalize_text(cell) for row in src_rows for cell in row]
    normalized_table_values = [normalize_text(cell) for row in out_rows for cell in row]

    if source_table_values != normalized_table_values:
        issues.append("Table content/order mismatch")

    # --------------------------------------------------------
    # Heading preservation
    # --------------------------------------------------------

    source_heading_texts = set()

    for paragraph in source_doc.paragraphs:
        style_name = paragraph.style.name if paragraph.style else ""

        if style_name.startswith("Heading"):
            text = normalize_text(get_paragraph_text(paragraph))

            if text:
                source_heading_texts.add(text)

    normalized_heading_texts = set()

    for hierarchy in normalized["hierarchies"]:
        normalized_heading_texts.update(normalize_text(x) for x in hierarchy)

    if normalized_doc.title:
        normalized_heading_texts.add(normalize_text(normalized_doc.title))

    missing_headings = source_heading_texts - normalized_heading_texts

    if missing_headings:
        issues.append(f"Missing heading content: {len(missing_headings)}")

    # --------------------------------------------------------
    # Full content + document order
    # --------------------------------------------------------

    src_sequence = source_atomic_sequence(source_doc, numbering_map)
    out_sequence = normalized_atomic_sequence(normalized_doc)

    if src_sequence != out_sequence:
        first_difference = None

        for i in range(min(len(src_sequence), len(out_sequence))):
            if src_sequence[i] != out_sequence[i]:
                first_difference = i
                break

        if first_difference is None and len(src_sequence) != len(out_sequence):
            first_difference = min(len(src_sequence), len(out_sequence))

        issues.append(
            f"Content/order mismatch: source={len(src_sequence)}, normalized={len(out_sequence)}, first_difference={first_difference}"
        )

    # --------------------------------------------------------
    # Duplicate warning only
    # --------------------------------------------------------

    block_texts = [normalize_text(x) for x in normalized["texts"] if normalize_text(x)]
    block_texts.extend(text for text, _, _ in normalized["lists"] if text)

    duplicate_counts = Counter(block_texts)
    suspicious_duplicates = {text: count for text, count in duplicate_counts.items() if count > 1 and len(text) > 30}

    return {
        "filename": os.path.basename(file_path),
        "issues": issues,
        "source_bullets": len(source_bullet_items),
        "source_decimal": len(source_decimal_items),
        "source_manual": len(source_manual_items),
        "source_hyperlinks": len(source_hyperlinks),
        "source_tables": len(source_tables),
        "normalized_tables": len(output_tables),
        "source_empty_cells": source_empty_cells,
        "normalized_empty_cells": normalized_empty_cells,
        "source_list_groups": len(source_groups),
        "normalized_list_groups": len(normalized_groups),
        "source_sequence": len(src_sequence),
        "normalized_sequence": len(out_sequence),
        "suspicious_duplicates": len(suspicious_duplicates),
    }


# ============================================================
# FULL VALIDATION
# ============================================================

def run_validation():
    print("=" * 72)
    print("STEP 7A: FINAL DOCX SEMANTIC VALIDATION")
    print("=" * 72)

    corpus_dir = os.path.abspath("Dataset/RAG-Multi-Corpus-main/RAG-Multi-Corpus-main/datasets")

    if not os.path.exists(corpus_dir):
        corpus_dir = os.path.abspath("Dataset/RAG-Multi-Corpus-main/RAG-Multi-Corpus-main/datasets")

    files = glob.glob(os.path.join(corpus_dir, "**", "*.docx"), recursive=True)

    print(f"\nTotal DOCX files discovered: {len(files)}")

    targets = [
        "About Aventro Motors.docx",
        "Adaptive Cruise Control (ACC).docx",
        "Adjusting the Seat Positions.docx",
        "employee_handbook.docx",
    ]

    print("\n" + "=" * 72)
    print("REPRESENTATIVE VALIDATION")
    print("=" * 72)

    for target in targets:
        file_path = next((f for f in files if target.lower() in os.path.basename(f).lower()), None)

        if not file_path:
            print(f"\n[SKIP] {target}")
            continue

        result = validate_single_document(file_path)

        print(f"\nDocument: {result['filename']}")
        print(f"Native bullets: {result['source_bullets']}")
        print(f"Native decimals: {result['source_decimal']}")
        print(f"Manual ordered: {result['source_manual']}")
        print(f"Hyperlinks: {result['source_hyperlinks']}")
        print(f"Tables: {result['source_tables']} -> {result['normalized_tables']}")
        print(f"Empty cells: {result['source_empty_cells']} -> {result['normalized_empty_cells']}")
        print(f"List groups: {result['source_list_groups']} -> {result['normalized_list_groups']}")
        print(f"Content sequence: {result['source_sequence']} -> {result['normalized_sequence']}")

        if result["issues"]:
            print("Status: FAIL")

            for issue in result["issues"]:
                print(f"  - {issue}")

        else:
            print("Status: PASS")

    print("\n" + "=" * 72)
    print("FULL CORPUS VALIDATION")
    print("=" * 72)

    totals = Counter()
    failed_docs = []
    duplicate_docs = []

    for file_path in files:
        try:
            result = validate_single_document(file_path)

            totals["bullets"] += result["source_bullets"]
            totals["decimals"] += result["source_decimal"]
            totals["manual"] += result["source_manual"]
            totals["hyperlinks"] += result["source_hyperlinks"]
            totals["source_tables"] += result["source_tables"]
            totals["normalized_tables"] += result["normalized_tables"]
            totals["source_empty"] += result["source_empty_cells"]
            totals["normalized_empty"] += result["normalized_empty_cells"]

            if result["issues"]:
                failed_docs.append((result["filename"], result["issues"]))

            if result["suspicious_duplicates"]:
                duplicate_docs.append((result["filename"], result["suspicious_duplicates"]))

        except Exception as e:
            failed_docs.append((os.path.basename(file_path), [str(e)]))

    passed = len(files) - len(failed_docs)

    print(f"\nTotal DOCX files: {len(files)}")
    print(f"Semantic validation passed: {passed}")
    print(f"Semantic validation failed: {len(failed_docs)}")

    print("\nStructure Totals:")
    print(f"Native bullet items: {totals['bullets']}")
    print(f"Native decimal items: {totals['decimals']}")
    print(f"Manual ordered items: {totals['manual']}")
    print(f"Hyperlinks: {totals['hyperlinks']}")

    print("\nTable Validation:")
    print(f"Source tables: {totals['source_tables']}")
    print(f"Normalized TableBlocks: {totals['normalized_tables']}")
    print(f"Source empty cells: {totals['source_empty']}")
    print(f"Normalized empty cells: {totals['normalized_empty']}")

    print(f"\nDocuments with duplicate warnings: {len(duplicate_docs)}")

    if failed_docs:
        print("\nFirst 20 Failed Documents:")

        for filename, issues in failed_docs[:20]:
            print(f"\n  {filename}")

            for issue in issues:
                print(f"    - {issue}")

    print("\n" + "=" * 72)

    if not failed_docs:
        print("DOCX FINAL SEMANTIC VALIDATION PASSED")
    else:
        print("DOCX FINAL SEMANTIC VALIDATION NEEDS REVIEW")

    print("=" * 72)


if __name__ == "__main__":
    run_validation()