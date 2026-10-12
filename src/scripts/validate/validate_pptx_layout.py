import os
import glob
import re
from collections import Counter

import pptx
from pptx.enum.shapes import PP_PLACEHOLDER

from src.parsers.pptx_parser import convert_pptx_to_normalized_document, get_paragraph_text, get_list_info, get_slide_title_shape
from src.schemas import TextBlock, TableBlock, ListBlock


MANUAL_NUMBER_RE = re.compile(r"^(\d+)\.\s+(.+)$")


# ============================================================
# BASIC HELPERS
# ============================================================

def normalize_text(text):
    return re.sub(r"\s+", " ", text).strip()


def get_shape_text(shape):
    if shape is None or not shape.has_text_frame: return ""
    return "\n".join(get_paragraph_text(p) for p in shape.text_frame.paragraphs if get_paragraph_text(p)).strip()


def flatten_list(block, depth=0):
    result = []

    for item in block.items:
        result.append((normalize_text(item.text), block.list_type, depth))
        for sub in item.sub_lists: result.extend(flatten_list(sub, depth + 1))

    return result


def collect_normalized_list_groups(section):
    return [flatten_list(block) for block in section.blocks if isinstance(block, ListBlock)]


def normalized_table_matrix(block):
    rows = []
    if block.headers: rows.append([normalize_text(x) for x in block.headers])
    rows.extend([[normalize_text(x) for x in row] for row in block.rows])
    return rows


def source_table_matrix(shape):
    return [[normalize_text(cell.text) for cell in row.cells] for row in shape.table.rows]


# ============================================================
# SOURCE LIST GROUPS
# ============================================================

def source_list_groups(slide, title_shape):
    groups = []
    ordered_shapes = sorted(slide.shapes, key=lambda s: (int(s.top), int(s.left)))

    for shape in ordered_shapes:
        if title_shape is not None and shape._element is title_shape._element: continue
        if getattr(shape, "has_table", False) or not shape.has_text_frame: continue

        buffered = []

        def flush():
            nonlocal buffered
            if not buffered: return

            base_level = buffered[0][1]
            groups.append([(normalize_text(text), list_type, max(0, level - base_level)) for text, level, list_type in buffered])
            buffered = []

        for paragraph in shape.text_frame.paragraphs:
            text = get_paragraph_text(paragraph)
            if not text: continue

            is_list, level, list_type = get_list_info(paragraph, shape, slide)

            if is_list:
                if buffered:
                    base_level = buffered[0][1]
                    if level <= base_level and buffered[-1][2] != list_type: flush()

                buffered.append((text, level, list_type))
                continue

            match = MANUAL_NUMBER_RE.match(text)

            if match:
                if buffered and buffered[-1][2] != "ordered": flush()
                buffered.append((match.group(2).strip(), 0, "ordered"))
                continue

            flush()

        flush()

    return groups


# ============================================================
# SOURCE / NORMALIZED ORDER
# ============================================================

def source_slide_sequence(slide, title_shape):
    sequence = []

    for shape in sorted(slide.shapes, key=lambda s: (int(s.top), int(s.left))):
        if title_shape is not None and shape._element is title_shape._element: continue

        if getattr(shape, "has_table", False):
            for row in shape.table.rows:
                for cell in row.cells: sequence.append(("table_cell", normalize_text(cell.text)))
            continue

        if not shape.has_text_frame: continue

        for paragraph in shape.text_frame.paragraphs:
            text = get_paragraph_text(paragraph)
            if not text: continue

            is_list, _, _ = get_list_info(paragraph, shape, slide)

            if is_list:
                sequence.append(("content", normalize_text(text)))
                continue

            match = MANUAL_NUMBER_RE.match(text)
            if match: sequence.append(("content", normalize_text(match.group(2))))
            else: sequence.append(("content", normalize_text(text))

            )

    return sequence


def normalized_slide_sequence(section):
    sequence = []

    for block in section.blocks:
        if isinstance(block, TextBlock): sequence.append(("content", normalize_text(block.text)))

        elif isinstance(block, ListBlock):
            for text, _, _ in flatten_list(block): sequence.append(("content", text))

        elif isinstance(block, TableBlock):
            for row in normalized_table_matrix(block):
                for cell in row: sequence.append(("table_cell", cell))

    return sequence


# ============================================================
# HYPERLINKS
# ============================================================

def source_hyperlink_texts(slide):
    values = []

    def inspect_text_frame(text_frame):
        for paragraph in text_frame.paragraphs:
            for run in paragraph.runs:
                try:
                    address = run.hyperlink.address
                except Exception:
                    address = None

                if address and normalize_text(run.text): values.append(normalize_text(run.text))

    for shape in slide.shapes:
        if getattr(shape, "has_table", False):
            for row in shape.table.rows:
                for cell in row.cells: inspect_text_frame(cell.text_frame)
        elif shape.has_text_frame:
            inspect_text_frame(shape.text_frame)

    return values


def normalized_slide_text(section):
    values = []

    values.extend(normalize_text(x) for x in section.hierarchy)

    for block in section.blocks:
        if isinstance(block, TextBlock): values.append(normalize_text(block.text))

        elif isinstance(block, ListBlock):
            values.extend(text for text, _, _ in flatten_list(block))

        elif isinstance(block, TableBlock):
            for row in normalized_table_matrix(block): values.extend(row)

    return " ".join(values)


# ============================================================
# SLIDE NUMBER VALIDATION
# ============================================================

def validate_block_slide_numbers(block, expected_slide):
    issues = []

    if getattr(block, "slide_number", None) != expected_slide: issues.append(f"{type(block).__name__} slide_number={getattr(block, 'slide_number', None)} expected={expected_slide}")

    if isinstance(block, ListBlock):
        for item in block.items:
            for sub in item.sub_lists: issues.extend(validate_block_slide_numbers(sub, expected_slide))

    return issues


# ============================================================
# SINGLE DOCUMENT
# ============================================================

def validate_document(file_path):
    prs = pptx.Presentation(file_path)
    normalized = convert_pptx_to_normalized_document(file_path)

    issues = []
    stats = Counter()

    if len(prs.slides) != len(normalized.sections): issues.append(f"Section count mismatch: slides={len(prs.slides)}, sections={len(normalized.sections)}")

    slide_count = min(len(prs.slides), len(normalized.sections))

    for slide_index in range(slide_count):
        slide = prs.slides[slide_index]
        section = normalized.sections[slide_index]
        slide_number = slide_index + 1
        title_shape = get_slide_title_shape(slide)
        source_title = normalize_text(get_shape_text(title_shape))
        normalized_title = normalize_text(section.hierarchy[0]) if section.hierarchy else ""

        # ----------------------------------------------------
        # TITLE
        # ----------------------------------------------------

        if source_title != normalized_title: issues.append(f"Slide {slide_number}: title mismatch")

        # ----------------------------------------------------
        # LIST SOURCE COUNTS
        # ----------------------------------------------------

        for shape in slide.shapes:
            if title_shape is not None and shape._element is title_shape._element: continue
            if getattr(shape, "has_table", False) or not shape.has_text_frame: continue

            for paragraph in shape.text_frame.paragraphs:
                text = get_paragraph_text(paragraph)
                if not text: continue

                is_list, _, list_type = get_list_info(paragraph, shape, slide)

                if is_list:
                    if list_type == "unordered": stats["native_bullets"] += 1
                    else: stats["native_numbered"] += 1
                    continue

                if MANUAL_NUMBER_RE.match(text): stats["manual_numbered"] += 1

        # ----------------------------------------------------
        # LIST STRUCTURE
        # ----------------------------------------------------

        source_groups = source_list_groups(slide, title_shape)
        output_groups = collect_normalized_list_groups(section)

        stats["source_list_groups"] += len(source_groups)
        stats["normalized_list_groups"] += len(output_groups)
        stats["source_list_items"] += sum(len(group) for group in source_groups)
        stats["normalized_list_items"] += sum(len(group) for group in output_groups)

        if source_groups != output_groups: issues.append(f"Slide {slide_number}: list nesting/type mismatch")

        # ----------------------------------------------------
        # TABLES
        # ----------------------------------------------------

        source_tables = [shape for shape in sorted(slide.shapes, key=lambda s: (int(s.top), int(s.left))) if getattr(shape, "has_table", False)]
        output_tables = [block for block in section.blocks if isinstance(block, TableBlock)]

        stats["source_tables"] += len(source_tables)
        stats["normalized_tables"] += len(output_tables)

        if len(source_tables) != len(output_tables):
            issues.append(f"Slide {slide_number}: table count mismatch {len(source_tables)} -> {len(output_tables)}")
        else:
            for table_index, (source_table, output_table) in enumerate(zip(source_tables, output_tables), 1):
                source_matrix = source_table_matrix(source_table)
                output_matrix = normalized_table_matrix(output_table)

                stats["source_empty_cells"] += sum(1 for row in source_matrix for cell in row if cell == "")
                stats["normalized_empty_cells"] += sum(1 for row in output_matrix for cell in row if cell == "")

                if source_matrix != output_matrix: issues.append(f"Slide {slide_number}: table {table_index} content/order mismatch")

        # ----------------------------------------------------
        # HYPERLINKS
        # ----------------------------------------------------

        source_links = source_hyperlink_texts(slide)
        output_text = normalized_slide_text(section)
        stats["hyperlinks"] += len(source_links)

        missing_links = [text for text in source_links if text not in output_text]
        if missing_links: issues.append(f"Slide {slide_number}: missing hyperlink visible text={len(missing_links)}")

        # ----------------------------------------------------
        # CONTENT ORDER
        # ----------------------------------------------------

        source_sequence = source_slide_sequence(slide, title_shape)
        output_sequence = normalized_slide_sequence(section)

        stats["source_sequence_items"] += len(source_sequence)
        stats["normalized_sequence_items"] += len(output_sequence)

        if source_sequence != output_sequence:
            first_difference = None

            for i in range(min(len(source_sequence), len(output_sequence))):
                if source_sequence[i] != output_sequence[i]:
                    first_difference = i
                    break

            if first_difference is None: first_difference = min(len(source_sequence), len(output_sequence))

            issues.append(f"Slide {slide_number}: content/order mismatch source={len(source_sequence)} normalized={len(output_sequence)} first_difference={first_difference}")

        # ----------------------------------------------------
        # SLIDE NUMBER
        # ----------------------------------------------------

        for block in section.blocks:
            slide_issues = validate_block_slide_numbers(block, slide_number)
            for issue in slide_issues: issues.append(f"Slide {slide_number}: {issue}")

    return {"filename": os.path.basename(file_path), "issues": issues, "stats": stats}


# ============================================================
# FULL CORPUS
# ============================================================

def main():
    print("=" * 72)
    print("STEP 7B: FINAL PPTX SEMANTIC VALIDATION")
    print("=" * 72)

    corpus = os.path.abspath("Dataset/RAG-Multi-Corpus-main/RAG-Multi-Corpus-main/datasets")
    if not os.path.exists(corpus): corpus = os.path.abspath("Dataset/RAG-Multi-Corpus-main/RAG-Multi-Corpus-main/datasets")

    files = glob.glob(os.path.join(corpus, "**", "*.pptx"), recursive=True)
    totals = Counter()
    failures = []

    targets = ["About Aventro Motors.pptx", "Adaptive Cruise Control (ACC).pptx", "Adjusting the Seat Positions.pptx", "employee_handbook.pptx"]

    print(f"\nTotal PPTX files discovered: {len(files)}")
    print("\n" + "=" * 72)
    print("REPRESENTATIVE VALIDATION")
    print("=" * 72)

    for target in targets:
        path = next((f for f in files if target.lower() == os.path.basename(f).lower()), None)

        if not path:
            print(f"\n[SKIP] {target}")
            continue

        result = validate_document(path)
        s = result["stats"]

        print(f"\nDocument: {result['filename']}")
        print(f"Native bullets: {s['native_bullets']}")
        print(f"Native numbered: {s['native_numbered']}")
        print(f"Manual numbered: {s['manual_numbered']}")
        print(f"List groups: {s['source_list_groups']} -> {s['normalized_list_groups']}")
        print(f"List items: {s['source_list_items']} -> {s['normalized_list_items']}")
        print(f"Tables: {s['source_tables']} -> {s['normalized_tables']}")
        print(f"Empty cells: {s['source_empty_cells']} -> {s['normalized_empty_cells']}")
        print(f"Hyperlinks: {s['hyperlinks']}")
        print(f"Content sequence: {s['source_sequence_items']} -> {s['normalized_sequence_items']}")

        if result["issues"]:
            print("Status: FAIL")
            for issue in result["issues"][:20]: print(f"  - {issue}")
        else:
            print("Status: PASS")

    print("\n" + "=" * 72)
    print("FULL CORPUS VALIDATION")
    print("=" * 72)

    for path in files:
        try:
            result = validate_document(path)
            totals.update(result["stats"])

            if result["issues"]: failures.append((result["filename"], result["issues"]))
        except Exception as e:
            failures.append((os.path.basename(path), [str(e)]))

    print(f"\nTotal PPTX files: {len(files)}")
    print(f"Semantic validation passed: {len(files) - len(failures)}")
    print(f"Semantic validation failed: {len(failures)}")

    print("\nList Validation:")
    print(f"Native bullet items: {totals['native_bullets']}")
    print(f"Native numbered items: {totals['native_numbered']}")
    print(f"Manual numbered items: {totals['manual_numbered']}")
    print(f"Source list groups: {totals['source_list_groups']}")
    print(f"Normalized list groups: {totals['normalized_list_groups']}")
    print(f"Source list items: {totals['source_list_items']}")
    print(f"Normalized list items: {totals['normalized_list_items']}")

    print("\nTable Validation:")
    print(f"Source tables: {totals['source_tables']}")
    print(f"Normalized TableBlocks: {totals['normalized_tables']}")
    print(f"Source empty cells: {totals['source_empty_cells']}")
    print(f"Normalized empty cells: {totals['normalized_empty_cells']}")

    print("\nOther Validation:")
    print(f"Hyperlinks checked: {totals['hyperlinks']}")
    print(f"Source content sequence items: {totals['source_sequence_items']}")
    print(f"Normalized sequence items: {totals['normalized_sequence_items']}")

    if failures:
        print("\nFirst 20 Failed Documents:")

        for filename, issues in failures[:20]:
            print(f"\n  {filename}")
            for issue in issues[:10]: print(f"    - {issue}")

    print("\n" + "=" * 72)
    print("PPTX FINAL SEMANTIC VALIDATION PASSED" if not failures else "PPTX FINAL SEMANTIC VALIDATION NEEDS REVIEW")
    print("=" * 72)


if __name__ == "__main__": main()